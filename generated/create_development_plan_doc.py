from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "docs/飞书自建应用研发样品管理系统_开发计划方案.docx"
LOGIC_DIAGRAM = "diagrams/飞书自建应用_研发样品管理平台_总体逻辑架构图_清晰版.png"
TECH_DIAGRAM = "diagrams/飞书自建应用_研发样品管理系统_技术系统架构图.png"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(89, 89, 89)
BLACK = RGBColor(0, 0, 0)
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
CALLOUT = "F4F6F9"


def set_run_font(run, size=None, bold=None, color=None, name="Microsoft YaHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, val in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9DEE7", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = width
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)


def set_paragraph_border_bottom(paragraph, color="2E74B5", size="12"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)


def add_para(doc, text="", style=None, size=11, bold=False, color=BLACK, before=0, after=6, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    if level == 1:
        size, color, before, after = 16, BLUE, 16, 8
    elif level == 2:
        size, color, before, after = 13, BLUE, 12, 6
    else:
        size, color, before, after = 12, DARK_BLUE, 8, 4
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=True, color=color)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        r = p.add_run(item)
        set_run_font(r, size=10.5)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        r = p.add_run(item)
        set_run_font(r, size=10.5)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table, color="DDE3EA", size="4")
    set_table_width(table, [Inches(6.45)])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    set_cell_margins(cell, top=130, bottom=130, start=160, end=160)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    set_run_font(r, size=10.5, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.2, color=BLACK)
    add_para(doc, "", after=2)


def add_table(doc, headers, rows, widths, fill=LIGHT_GRAY, font_size=8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_borders(table)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, fill)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_run_font(r, size=font_size, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if idx == 0 or len(str(value)) <= 12:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    set_table_width(table, widths)
    add_para(doc, "", after=2)
    return table


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table, color="DDE3EA", size="4")
    set_table_width(table, [Inches(6.45)])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F8FAFC")
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("\n".join(lines))
    set_run_font(r, name="Consolas", size=8.6, color=RGBColor(30, 41, 59))
    add_para(doc, "", after=2)


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DARK_BLUE),
    ]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True

    header = section.header.paragraphs[0]
    header.text = "飞书自建应用研发样品管理系统开发计划方案"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run, size=9, color=GRAY)

    footer = section.footer.paragraphs[0]
    footer.text = "内部规划稿 | v1.0"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        set_run_font(run, size=9, color=GRAY)
    return doc


doc = setup_doc()

add_para(doc, "内部开发计划方案", size=11, bold=True, color=GRAY, after=8)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
r = p.add_run("飞书自建应用研发样品管理系统")
set_run_font(r, size=24, bold=True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
r = p.add_run("完整流程版：样品需求、研发打样、测试、寄样、核价、转订单、工艺/BOM 与归档")
set_run_font(r, size=13.5, color=GRAY)
for label, value in [
    ("方案版本", "v1.0"),
    ("后端技术", "Java 17 + Spring Boot 3"),
    ("前端技术", "Vue3 + Vant / Ant Design Vue"),
    ("数据与文件", "PostgreSQL + Redis + MinIO/NAS"),
    ("协同平台", "飞书自建应用、消息卡片、审批确认、云盘文件"),
    ("文档目的", "为研发样品管理系统立项、开发排期、接口设计和验收提供统一依据"),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    rr = p.add_run(f"{label}：")
    set_run_font(rr, size=10.5, bold=True)
    rv = p.add_run(value)
    set_run_font(rv, size=10.5)
rule = doc.add_paragraph()
set_paragraph_border_bottom(rule)

add_callout(
    doc,
    "建议结论",
    "采用“飞书做入口与协同，自研系统做主数据与流程”的建设方式。首版按完整业务链路规划，但建议分阶段上线：先打通研发样品闭环，再扩展寄样反馈、财务报价、转订单、工艺文件和BOM任务。"
)

add_heading(doc, "1. 项目背景与建设目标", 1)
add_para(doc, "当前研发样品从需求提出、研发分发、现场打样、测试、寄样、客户反馈到核价和量产资料准备之间，存在纸质记录、文件散落、版本不清、沟通依赖人工提醒等问题。系统建设目标是把样品全生命周期沉淀到一套可追踪、可版本化、可归档的流程中。")
add_bullets(doc, [
    "统一入口：所有角色通过飞书工作台进入研发样品管理应用。",
    "现场录入：研发人员用手机/平板在实验现场录入打样实验单、称重数据和过程照片。",
    "流程闭环：覆盖样品需求、任务分发、测试、寄样、反馈、核价、转订单、工艺/BOM任务。",
    "版本管控：实验单、样品版本、核价文件、工艺文件、BOM任务均保留版本，不覆盖历史。",
    "文件归档：产品档案按产品和版本归档，便于追溯与复盘。",
])

add_heading(doc, "2. 建设范围", 1)
add_table(doc, ["范围", "包含内容", "说明"], [
    ["首版必须包含", "样品需求、任务分发、研发接受、打样实验单、测试结果、Excel核价文件、归档、飞书通知", "形成可运行闭环"],
    ["首版规划包含", "寄样登记、客户反馈、财务报价状态、转待下单、工艺文件任务、BOM录入任务", "可按二期节奏上线"],
    ["暂不深做", "完整财务成本系统、ERP生产计划、电子秤自动采集、标签营养成分合规校验", "预留接口和字段"],
], [Inches(1.35), Inches(3.05), Inches(2.05)], font_size=9)

add_heading(doc, "3. 总体逻辑架构", 1)
add_para(doc, "总体架构采用七层逻辑：用户层、飞书统一入口层、平台控制层、业务服务层、工具模板层、数据归档层和基础设施层。")
doc.add_picture(LOGIC_DIAGRAM, width=Inches(6.45))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_para(doc, "图 1：飞书自建应用研发样品管理平台总体逻辑架构图", size=9.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)

add_heading(doc, "4. 技术系统架构", 1)
add_para(doc, "技术架构采用前后端分离。飞书自建应用作为入口，Nginx 负责 HTTPS 和反向代理，Spring Boot 负责业务服务，PostgreSQL 负责业务数据，Redis 负责缓存和异步任务协调，MinIO 或 NAS 负责附件、Excel、工艺文件和BOM文件归档。")
doc.add_picture(TECH_DIAGRAM, width=Inches(6.45))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_para(doc, "图 2：研发样品管理飞书自建应用技术系统架构图", size=9.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)

add_heading(doc, "5. 推荐技术栈", 1)
add_table(doc, ["层级", "技术选型", "用途"], [
    ["前端移动端", "Vue3 + Vant", "手机/平板现场录入、拍照上传、待办处理"],
    ["前端PC端", "Vue3 + Ant Design Vue", "研发总监、内勤、财务、管理员使用的管理端"],
    ["后端", "Java 17 + Spring Boot 3", "REST API、状态机、权限、业务流程服务"],
    ["数据库", "PostgreSQL", "样品、任务、实验单、版本、文件索引、审计日志"],
    ["缓存/队列", "Redis", "登录状态缓存、异步任务、消息通知队列"],
    ["文件存储", "MinIO 优先，NAS 可选", "照片、附件、Excel核价文件、工艺文件、BOM文件"],
    ["Excel生成", "EasyExcel + Apache POI", "按现有核价表模板生成Excel文件"],
    ["部署", "Nginx + Docker Compose", "内网或云主机部署，便于维护和迁移"],
    ["飞书", "自建应用 + OAuth + 消息卡片 + 审批 + 云盘", "统一入口、通知、确认和文件协同"],
], [Inches(1.25), Inches(2.25), Inches(2.95)], font_size=8.8)

add_heading(doc, "6. 角色与权限", 1)
add_table(doc, ["角色", "主要操作权限", "关键限制"], [
    ["销售内勤", "录入样品需求、补充需求、寄样登记、录入客户反馈", "不能修改已提交实验单和测试结论"],
    ["研发总监", "审核需求、退回补充、分发任务、转派任务、停止确认、查看全量进度", "不能覆盖历史版本"],
    ["研发人员", "接受任务、填写打样实验单、上传照片、提交实验版本", "提交后只能新建下一版"],
    ["测试/品控", "查看待测试样品、填写测试结果、上传测试附件", "不能修改研发配方"],
    ["财务", "查看/下载核价文件、维护报价状态、上传报价附件", "不参与实验单编辑"],
    ["业务员/管理层", "查看样品进度、客户反馈、报价状态和转下单结果", "默认只读"],
    ["管理员", "维护人员权限、工艺模板、物料库、Excel模板、系统参数", "需要记录审计日志"],
], [Inches(1.35), Inches(3.65), Inches(1.45)], font_size=8.6)

add_heading(doc, "7. 核心业务流程", 1)
add_numbered(doc, [
    "销售内勤在飞书应用中录入样品需求，填写客户、产品、规格、样品数量、期望日期、包装要求和备注。",
    "研发总监审核需求，资料不完整则退回销售内勤补充；资料完整则分发给指定研发人员。",
    "研发人员在飞书消息卡片或应用待办中接受任务，进入打样阶段。",
    "研发人员在现场使用手机/平板填写打样实验单，录入原辅料、包材、工序称重、损耗、得率、照片和附件。",
    "研发提交实验单后，系统锁定当前样品版本，通知测试/品控人员进行测试。",
    "测试人员提交测试结果：通过、需要复打样或建议停止。",
    "测试通过后样品完成，系统通知销售内勤寄样，并同步生成Excel核价基础文件。",
    "销售内勤登记快递单号、寄样数量和样品版本，等待业务员或客户反馈。",
    "财务下载核价文件并维护报价状态，报价结果反馈给业务员。",
    "客户通过后转入待下单订单，并生成计划制作工艺文件和产品BOM录入任务；客户不通过则进入下一版打样；项目取消则进入停止/废弃项目池。",
])

add_heading(doc, "8. 状态机设计", 1)
add_code_block(doc, [
    "待审核需求 -> 退回补充 -> 待审核需求",
    "待审核需求 -> 待分发 -> 待研发接受 -> 打样中 -> 待测试",
    "待测试 -> 测试不通过需复打样 -> 待研发接受",
    "待测试 -> 样品完成 -> 待寄样 -> 已寄样待反馈",
    "样品完成 -> 待财务核价 -> 已报价",
    "已寄样待反馈 -> 客户不通过需复打样 -> 待研发接受",
    "已报价 + 客户通过 -> 待下单 -> 工艺BOM准备 -> 已归档",
    "待测试/已寄样待反馈 -> 已停止 -> 停止/废弃项目池",
])
add_para(doc, "状态流转必须由后端状态机统一控制。前端不能直接写状态字段，所有状态变更都应通过明确动作触发，并写入审计日志。")

add_heading(doc, "9. 版本规则", 1)
add_table(doc, ["对象", "版本示例", "规则"], [
    ["样品版本", "A0、A1、A2", "每次复打样生成下一版，历史不可覆盖"],
    ["实验单版本", "跟随样品版本", "提交后锁定，只允许查看和复制生成下一版"],
    ["核价文件版本", "A0-核价V1、A0-核价V2、A1-核价V1", "同一样品版本可多次生成核价文件版本"],
    ["工艺文件版本", "A1-工艺V1", "客户通过后生成，后续修改独立升版"],
    ["产品BOM版本", "A1-BOM-V1", "与通过的样品版本关联"],
], [Inches(1.6), Inches(2.0), Inches(2.85)], font_size=8.8)

add_heading(doc, "10. 数据模型设计", 1)
add_table(doc, ["数据表", "中文含义", "核心字段"], [
    ["sample_request", "样品需求单", "sample_no, product_name, customer_name, specification, required_date, status"],
    ["rnd_task", "研发任务", "task_no, request_id, assignee_id, director_id, status, assigned_at"],
    ["sample_version", "样品版本", "sample_no, version_no, stage, status, locked_at, submitted_at"],
    ["experiment_form", "打样实验单", "version_id, product_type, formula_summary, yield_rate, conclusion"],
    ["experiment_material", "原辅料/包材明细", "form_id, process_name, ns_code, material_name, weight_kg, utilization_rate"],
    ["experiment_process", "工序称重记录", "form_id, process_name, before_weight, after_weight, loss_rate"],
    ["test_record", "测试记录", "version_id, tester_id, result, comments, attachment_ids"],
    ["shipment_record", "寄样记录", "version_id, express_no, recipient, sample_qty, shipped_at"],
    ["customer_feedback", "客户反馈", "shipment_id, result, feedback_detail, feedback_at"],
    ["pricing_file", "核价文件", "version_id, pricing_version, file_path, feishu_file_url, status"],
    ["finance_quote", "财务报价状态", "pricing_file_id, quote_status, quote_amount, attachment_id"],
    ["process_file_task", "工艺文件任务", "version_id, process_version, owner_id, status"],
    ["product_bom_task", "产品BOM任务", "version_id, bom_version, owner_id, status"],
    ["archive_file", "归档文件", "biz_type, biz_id, file_name, storage_path, feishu_file_url"],
    ["audit_log", "审计日志", "operator_id, action, before_state, after_state, operated_at"],
], [Inches(1.55), Inches(1.35), Inches(3.55)], font_size=7.8)

add_heading(doc, "11. API设计", 1)
add_para(doc, "后端 REST API 统一使用 `/api/v1` 前缀。飞书回调接口独立使用 `/feishu` 前缀。")
add_code_block(doc, [
    "POST   /api/v1/sample-requests",
    "GET    /api/v1/sample-requests",
    "POST   /api/v1/sample-requests/{id}/submit",
    "POST   /api/v1/sample-requests/{id}/return",
    "POST   /api/v1/rnd-tasks/{id}/assign",
    "POST   /api/v1/rnd-tasks/{id}/accept",
    "POST   /api/v1/rnd-tasks/{id}/transfer",
    "POST   /api/v1/sample-versions/{id}/experiment-form",
    "POST   /api/v1/sample-versions/{id}/submit-experiment",
    "GET    /api/v1/sample-versions/{id}/archive",
    "POST   /api/v1/sample-versions/{id}/test-records",
    "POST   /api/v1/sample-versions/{id}/request-resample",
    "POST   /api/v1/sample-versions/{id}/stop",
    "POST   /api/v1/sample-versions/{id}/shipments",
    "POST   /api/v1/shipments/{id}/feedback",
    "POST   /api/v1/sample-versions/{id}/pricing-files",
    "POST   /api/v1/pricing-files/{id}/submit-to-finance",
    "POST   /api/v1/pricing-files/{id}/finance-quote",
    "POST   /api/v1/sample-versions/{id}/convert-to-order",
    "POST   /api/v1/sample-versions/{id}/process-file-task",
    "POST   /api/v1/sample-versions/{id}/bom-task",
    "GET    /feishu/oauth/callback",
    "POST   /feishu/events",
    "POST   /feishu/card-actions",
])

add_heading(doc, "12. 飞书集成设计", 1)
add_table(doc, ["能力", "使用方式", "业务场景"], [
    ["飞书免登/OAuth", "前端获取授权码，后端换取用户身份并映射系统账号", "用户从飞书工作台进入应用，无需重复登录"],
    ["飞书消息卡片", "后端调用飞书消息接口推送卡片，卡片按钮跳转或回调", "任务分发、待测试、待寄样、核价文件生成提醒"],
    ["飞书审批/确认", "关键节点可接入飞书审批，也可使用系统内确认按钮", "需求审核、测试确认、停止打样确认"],
    ["飞书云盘", "系统生成文件后上传云盘，保存文件链接", "核价文件、测试附件、归档文件共享"],
    ["多维表格", "作为可选看板，不作为主数据库", "样品进度、任务量、逾期统计展示"],
], [Inches(1.35), Inches(3.0), Inches(2.1)], font_size=8.7)

add_heading(doc, "13. Excel核价文件设计", 1)
add_para(doc, "首版以现有核价清单样表为基础生成单 Sheet Excel 文件。该文件定位为研发核价基础文件，供财务进一步核算报价使用。")
add_bullets(doc, [
    "文件命名：产品名称-核价原料清单-样品版本-核价版本.xlsx。",
    "基础信息：公司名称、产品名称、负责人、规格、类别、状态、版本号、文件编号、实施日期、编写人。",
    "原料清单：工序分类、编号、NS编码、物料名称、重量kg、参考利用率、领料重量kg、备注。",
    "包装物料：编号、NS编码、物料名称、内部代码、数量、包装规格、备注。",
    "研发结果：研发参考出成kg、得率、参考包数。",
    "归档规则：生成后保存到文件库，生成 pricing_file 记录，并可同步飞书云盘链接。",
])
add_code_block(doc, [
    "领料重量 = 重量 / 参考利用率",
    "参考包数 = 研发参考出成kg / 单袋规格kg",
    "得率 = 研发参考出成kg / 关键原料领料重量",
])

add_heading(doc, "14. 文件归档设计", 1)
add_code_block(doc, [
    "产品档案/",
    "  500g香卤大肠头/",
    "    A0/",
    "      样品需求单.json",
    "      打样实验单.json",
    "      现场照片/",
    "      测试记录/",
    "      核价文件_A0_核价V1.xlsx",
    "      寄样记录.json",
    "      客户反馈.json",
    "    A1/",
    "      打样实验单.json",
    "      核价文件_A1_核价V1.xlsx",
    "      工艺文件_A1_工艺V1.docx",
    "      产品BOM_A1_BOM_V1.xlsx",
])
add_para(doc, "文件归档以数据库索引为准，目录结构用于运维和人工排查。所有业务页面应通过 archive_file 记录读取文件，而不是直接拼接文件路径。")

add_heading(doc, "15. 前端页面规划", 1)
add_table(doc, ["页面", "使用角色", "核心功能"], [
    ["首页看板", "全部角色", "我的待办、样品进度、逾期提醒、快速入口"],
    ["样品需求", "销售内勤、研发总监", "新建、补充、审核、退回、查看需求"],
    ["研发任务", "研发总监、研发人员", "任务分发、接受、转派、进度跟踪"],
    ["打样实验单", "研发人员", "配方、工序、称重、损耗、得率、照片附件"],
    ["测试确认", "测试/品控", "测试结果、问题点、复打样或通过结论"],
    ["寄样反馈", "销售内勤、业务员", "快递登记、客户反馈、通过/调整/停止"],
    ["核价文件", "研发、财务", "生成Excel、版本列表、提交财务、报价状态"],
    ["产品档案", "管理层、研发总监", "按产品和版本查看全量资料"],
    ["基础配置", "管理员", "人员权限、产品分类、工艺模板、Excel模板"],
], [Inches(1.45), Inches(1.75), Inches(3.25)], font_size=8.6)

add_heading(doc, "16. 实施排期建议", 1)
add_table(doc, ["阶段", "周期", "交付内容", "验收标准"], [
    ["阶段1：项目骨架", "1-2周", "Spring Boot、Vue、数据库、登录框架、Docker部署骨架", "本地和测试环境可启动，能通过飞书进入系统"],
    ["阶段2：研发闭环", "3-5周", "需求录入、任务分发、研发接受、打样实验单、版本锁定、测试确认", "一个样品能从需求走到样品完成"],
    ["阶段3：核价与归档", "2-3周", "Excel核价文件生成、文件上传、产品档案、飞书通知", "能按样表导出核价文件并归档"],
    ["阶段4：寄样与反馈", "2-3周", "寄样登记、客户反馈、复打样循环、停止/废弃项目池", "客户不通过能自动回到下一版打样"],
    ["阶段5：财务报价与转单", "2-3周", "财务报价状态、待下单、工艺文件任务、BOM任务", "客户通过后能生成后续任务"],
    ["阶段6：试运行优化", "2-4周", "权限、看板、报表、性能、备份、操作培训", "至少选取2-3个真实产品完整试跑"],
], [Inches(1.45), Inches(0.85), Inches(2.75), Inches(1.4)], font_size=8.1)

add_heading(doc, "17. 测试与验收", 1)
add_heading(doc, "17.1 单元测试", 2)
add_bullets(doc, [
    "状态机流转：合法状态可流转，非法状态拒绝。",
    "版本锁定：提交后不可修改，复打样生成新版本。",
    "Excel生成：字段填充、公式计算、文件名、版本号正确。",
    "权限控制：不同角色只能访问允许接口。",
    "飞书用户映射：feishu_user_id 能正确关联系统用户。",
])
add_heading(doc, "17.2 集成测试", 2)
add_bullets(doc, [
    "销售内勤提交需求 -> 研发总监分发 -> 研发接受任务。",
    "研发提交实验单 -> 测试人员收到通知 -> 测试通过。",
    "测试不通过 -> 自动生成下一版打样任务。",
    "测试通过 -> 样品完成 -> 生成核价文件 -> 财务收到通知。",
    "寄样登记 -> 客户反馈通过 -> 转待下单 -> 生成工艺/BOM任务。",
    "停止打样 -> 进入停止/废弃项目池 -> 文件归档完整。",
])
add_heading(doc, "17.3 业务验收", 2)
add_bullets(doc, [
    "手机/平板能完成现场打样录入和拍照上传。",
    "Excel导出格式接近现有核价清单模板。",
    "任意产品档案能按版本查看实验单、图片、测试、寄样、核价文件。",
    "飞书消息卡片可直接跳转对应任务。",
    "关键记录都有审计日志：创建人、修改人、提交人、锁定时间。",
])

add_heading(doc, "18. 风险与应对", 1)
add_table(doc, ["风险", "影响", "应对策略"], [
    ["飞书接口权限申请不完整", "免登、消息、云盘功能受阻", "提前由企业管理员创建自建应用并确认权限清单"],
    ["移动端外网访问公司内网系统困难", "现场无法通过飞书移动端访问", "配置HTTPS域名、VPN或安全网关，必要时采用云主机部署"],
    ["Excel模板频繁变化", "导出格式反复调整", "把模板做成可配置版本，先锁定首版样表"],
    ["研发现场录入不完整", "得率和核价文件不可信", "设置必填字段、草稿校验、提交前检查清单"],
    ["历史版本被覆盖", "追溯失效", "后端强制锁定提交版本，所有修改生成新版本"],
    ["文件散落在飞书和本地", "档案不完整", "以系统归档库为主，飞书云盘只作为协同链接"],
], [Inches(1.8), Inches(1.8), Inches(2.85)], font_size=8.4)

add_heading(doc, "19. 首批准备事项", 1)
add_bullets(doc, [
    "确认飞书企业管理员，创建“研发样品管理”自建应用。",
    "确认首版角色人员名单和权限边界。",
    "确认样品需求单字段、打样实验单字段和测试记录字段。",
    "确认Excel核价文件首版模板，以现有原料清单样表为基准。",
    "确认部署方式：公司内网主机、云主机、是否需要外网HTTPS域名。",
    "选取2-3个真实产品作为试点，例如冷冻即热菜、腊制品、生制调理品各一个。",
])

doc.save(OUT)
print(OUT)
